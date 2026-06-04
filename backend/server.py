from dotenv import load_dotenv
from pathlib import Path

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

import os
import logging
import secrets
import bcrypt
import jwt
import base64
from datetime import datetime, timezone, timedelta
from typing import List, Optional
from bson import ObjectId

from fastapi import FastAPI, APIRouter, HTTPException, Request, Response, Depends, UploadFile, File
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, EmailStr

# ---------- Setup ----------
mongo_url = os.environ['MONGO_URL']
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ['DB_NAME']]

JWT_ALGORITHM = "HS256"
JWT_SECRET = os.environ['JWT_SECRET']
ADMIN_EMAIL = os.environ.get('ADMIN_EMAIL', 'admin@nmp.com')
ADMIN_PASSWORD = os.environ.get('ADMIN_PASSWORD', 'admin123')

logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

app = FastAPI(title="NMP Sales OS")
api = APIRouter(prefix="/api")

# ---------- Helpers ----------
def hash_password(pw: str) -> str:
    return bcrypt.hashpw(pw.encode("utf-8"), bcrypt.gensalt()).decode("utf-8")

def verify_password(pw: str, hashed: str) -> bool:
    try:
        return bcrypt.checkpw(pw.encode("utf-8"), hashed.encode("utf-8"))
    except Exception:
        return False

def create_access_token(user_id: str, email: str, role: str) -> str:
    payload = {"sub": user_id, "email": email, "role": role,
               "exp": datetime.now(timezone.utc) + timedelta(hours=24), "type": "access"}
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)

def serialize(doc: dict) -> dict:
    if not doc:
        return doc
    doc = dict(doc)
    if "_id" in doc:
        doc["id"] = str(doc["_id"])
        del doc["_id"]
    doc.pop("password_hash", None)
    return doc

async def get_current_user(request: Request) -> dict:
    token = request.cookies.get("access_token")
    if not token:
        auth = request.headers.get("Authorization", "")
        if auth.startswith("Bearer "):
            token = auth[7:]
    if not token:
        raise HTTPException(401, "Not authenticated")
    try:
        payload = jwt.decode(token, JWT_SECRET, algorithms=[JWT_ALGORITHM])
        user = await db.users.find_one({"_id": ObjectId(payload["sub"])})
        if not user:
            raise HTTPException(401, "User not found")
        # force-logout via token_version mismatch
        if payload.get("tv", 0) != user.get("token_version", 0):
            raise HTTPException(401, "Session revoked")
        if user.get("status") in ("Suspended", "Terminated", "Inactive"):
            raise HTTPException(403, f"Account is {user.get('status').lower()}")
        return serialize(user)
    except jwt.ExpiredSignatureError:
        raise HTTPException(401, "Token expired")
    except jwt.InvalidTokenError:
        raise HTTPException(401, "Invalid token")

def require_roles(*roles):
    async def checker(user: dict = Depends(get_current_user)):
        if user.get("role") not in roles:
            raise HTTPException(403, "Insufficient permissions")
        return user
    return checker

# ---------- Models ----------
class LoginIn(BaseModel):
    email: EmailStr
    password: str

class SetupAdminIn(BaseModel):
    email: EmailStr
    password: str
    name: str

class EmployeeIn(BaseModel):
    name: str
    email: EmailStr
    mobile: Optional[str] = ""
    department: Optional[str] = ""
    designation: Optional[str] = ""
    territory: Optional[str] = ""
    role: str = "field_sales"
    reporting_manager: Optional[str] = ""

class InviteAcceptIn(BaseModel):
    token: str
    password: str

class LeadIn(BaseModel):
    lead_name: str
    company_name: Optional[str] = ""
    contact_person: Optional[str] = ""
    mobile: Optional[str] = ""
    email: Optional[str] = ""
    address: Optional[str] = ""
    city: Optional[str] = ""
    state: Optional[str] = ""
    lead_source: str = "Website"
    lead_type: str = "Product Sale"
    visit_requirement: str = "Not Required"
    expected_deal_value: float = 0
    priority: str = "Medium"
    assigned_to: Optional[str] = ""
    status: str = "New"
    notes: Optional[str] = ""
    next_follow_up: Optional[str] = ""
    follow_up_type: Optional[str] = "Call"
    decision_maker: Optional[str] = ""
    budget: float = 0
    expected_closure_date: Optional[str] = ""
    competitor: Optional[str] = ""
    probability: int = 25
    closure_type: Optional[str] = ""

class SiteVisitIn(BaseModel):
    gps_lat: float
    gps_lng: float
    gps_accuracy: Optional[float] = 0
    site_address: Optional[str] = ""
    requirement_notes: Optional[str] = ""
    photos: List[str] = []
    videos: List[str] = []

class TaskIn(BaseModel):
    lead_id: Optional[str] = ""
    title: str
    description: Optional[str] = ""
    task_type: str = "Follow-Up"
    assigned_to: Optional[str] = ""
    due_date: Optional[str] = ""
    status: str = "Open"

class QuotationItem(BaseModel):
    sku: Optional[str] = ""
    product_name: str
    quantity: int = 1
    unit_price: float = 0
    discount_pct: float = 0
    gst_rate: float = 18

class QuotationIn(BaseModel):
    lead_id: str
    quotation_number: Optional[str] = ""
    items: List[QuotationItem]
    installation_charge: float = 0
    freight_charge: float = 0
    amc_charge: float = 0
    misc_charge: float = 0
    terms: Optional[str] = "50% advance, balance against delivery. Warranty as per brand."
    notes: Optional[str] = ""
    status: Optional[str] = "Draft"

class InventoryItem(BaseModel):
    sku: str
    product_name: str
    brand: Optional[str] = ""
    category: Optional[str] = ""
    cost_price: float = 0
    selling_price: float = 0
    mrp: float = 0
    gst_rate: float = 18
    stock: int = 0
    warehouse: Optional[str] = ""

class ActivityIn(BaseModel):
    lead_id: str
    activity_type: str
    description: str

# ---------- Auth ----------
@api.get("/auth/setup-status")
async def setup_status():
    count = await db.users.count_documents({"role": "super_admin"})
    return {"setup_complete": count > 0}

@api.post("/auth/setup-admin")
async def setup_admin(body: SetupAdminIn, response: Response):
    if await db.users.count_documents({"role": "super_admin"}) > 0:
        raise HTTPException(400, "Super admin already exists")
    user = {
        "email": body.email.lower(),
        "password_hash": hash_password(body.password),
        "name": body.name,
        "role": "super_admin",
        "status": "Active",
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    res = await db.users.insert_one(user)
    token = create_access_token(str(res.inserted_id), user["email"], "super_admin")
    response.set_cookie("access_token", token, httponly=True, secure=True, samesite="none", max_age=86400, path="/")
    user["id"] = str(res.inserted_id)
    return {"user": serialize(user), "token": token}

@api.post("/auth/login")
async def login(body: LoginIn, request: Request, response: Response):
    user = await db.users.find_one({"email": body.email.lower()})
    if not user or not verify_password(body.password, user["password_hash"]):
        await db.login_history.insert_one({
            "email": body.email.lower(), "success": False,
            "ip": request.client.host if request.client else "",
            "user_agent": request.headers.get("User-Agent", ""),
            "at": datetime.now(timezone.utc).isoformat(),
        })
        raise HTTPException(401, "Invalid email or password")
    if user.get("status") in ("Inactive", "Suspended", "Terminated"):
        raise HTTPException(403, f"Account is {user.get('status', 'inactive').lower()}")
    tv = user.get("token_version", 0)
    payload = {"sub": str(user["_id"]), "email": user["email"], "role": user["role"], "tv": tv,
               "exp": datetime.now(timezone.utc) + timedelta(hours=24), "type": "access"}
    token = jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)
    response.set_cookie("access_token", token, httponly=True, secure=True, samesite="none", max_age=86400, path="/")
    await db.login_history.insert_one({
        "user_id": str(user["_id"]), "email": user["email"], "success": True,
        "ip": request.client.host if request.client else "",
        "user_agent": request.headers.get("User-Agent", ""),
        "at": datetime.now(timezone.utc).isoformat(),
    })
    return {"user": serialize(user), "token": token}

@api.post("/auth/logout")
async def logout(response: Response):
    response.delete_cookie("access_token", path="/")
    return {"ok": True}

@api.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return user

# ---------- Employees ----------
@api.get("/employees")
async def list_employees(user: dict = Depends(get_current_user)):
    users = await db.users.find().sort("created_at", -1).to_list(500)
    return [serialize(u) for u in users]

@api.post("/employees")
async def create_employee(body: EmployeeIn, user: dict = Depends(require_roles("super_admin", "admin"))):
    if await db.users.find_one({"email": body.email.lower()}):
        raise HTTPException(400, "Email already exists")
    invite_token = secrets.token_urlsafe(24)
    temp_pw = secrets.token_urlsafe(8)
    doc = {
        **body.model_dump(),
        "email": body.email.lower(),
        "password_hash": hash_password(temp_pw),
        "status": "Pending",
        "invite_token": invite_token,
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    res = await db.users.insert_one(doc)
    doc["id"] = str(res.inserted_id)
    return {"employee": serialize(doc), "invite_token": invite_token, "temp_password": temp_pw}

@api.put("/employees/{employee_id}")
async def update_employee(employee_id: str, body: dict, user: dict = Depends(require_roles("super_admin", "admin"))):
    body.pop("password_hash", None); body.pop("_id", None); body.pop("id", None)
    await db.users.update_one({"_id": ObjectId(employee_id)}, {"$set": body})
    u = await db.users.find_one({"_id": ObjectId(employee_id)})
    return serialize(u)

@api.post("/employees/invite-accept")
async def invite_accept(body: InviteAcceptIn, response: Response):
    u = await db.users.find_one({"invite_token": body.token})
    if not u:
        raise HTTPException(400, "Invalid invite token")
    await db.users.update_one(
        {"_id": u["_id"]},
        {"$set": {"password_hash": hash_password(body.password), "status": "Active"},
         "$unset": {"invite_token": ""}},
    )
    token = create_access_token(str(u["_id"]), u["email"], u["role"])
    response.set_cookie("access_token", token, httponly=True, secure=True, samesite="none", max_age=86400, path="/")
    u = await db.users.find_one({"_id": u["_id"]})
    return {"user": serialize(u), "token": token}

# ---------- Employee Admin Actions ----------
@api.post("/employees/{emp_id}/resend-invite")
async def resend_invite(emp_id: str, user: dict = Depends(require_roles("super_admin", "admin"))):
    new_token = secrets.token_urlsafe(24)
    res = await db.users.update_one({"_id": ObjectId(emp_id)}, {"$set": {"invite_token": new_token, "status": "Pending"}})
    if not res.matched_count:
        raise HTTPException(404, "Employee not found")
    return {"invite_token": new_token}

@api.post("/employees/{emp_id}/revoke-invite")
async def revoke_invite(emp_id: str, user: dict = Depends(require_roles("super_admin", "admin"))):
    await db.users.update_one({"_id": ObjectId(emp_id)}, {"$unset": {"invite_token": ""}, "$set": {"status": "Inactive"}})
    return {"ok": True}

@api.post("/employees/{emp_id}/suspend")
async def suspend_employee(emp_id: str, user: dict = Depends(require_roles("super_admin", "admin"))):
    await db.users.update_one({"_id": ObjectId(emp_id)}, {"$set": {"status": "Suspended"}, "$inc": {"token_version": 1}})
    return {"ok": True}

@api.post("/employees/{emp_id}/terminate")
async def terminate_employee(emp_id: str, user: dict = Depends(require_roles("super_admin", "admin"))):
    await db.users.update_one({"_id": ObjectId(emp_id)}, {"$set": {"status": "Terminated"}, "$inc": {"token_version": 1}})
    return {"ok": True}

@api.post("/employees/{emp_id}/activate")
async def activate_employee(emp_id: str, user: dict = Depends(require_roles("super_admin", "admin"))):
    await db.users.update_one({"_id": ObjectId(emp_id)}, {"$set": {"status": "Active"}})
    return {"ok": True}

@api.post("/employees/{emp_id}/reset-password")
async def reset_password(emp_id: str, user: dict = Depends(require_roles("super_admin", "admin"))):
    temp_pw = secrets.token_urlsafe(8)
    await db.users.update_one({"_id": ObjectId(emp_id)},
                              {"$set": {"password_hash": hash_password(temp_pw)},
                               "$inc": {"token_version": 1}})
    return {"temp_password": temp_pw}

@api.post("/employees/{emp_id}/force-logout")
async def force_logout(emp_id: str, user: dict = Depends(require_roles("super_admin", "admin"))):
    await db.users.update_one({"_id": ObjectId(emp_id)}, {"$inc": {"token_version": 1}})
    return {"ok": True}

@api.get("/employees/{emp_id}/login-history")
async def login_history(emp_id: str, user: dict = Depends(require_roles("super_admin", "admin"))):
    u = await db.users.find_one({"_id": ObjectId(emp_id)})
    if not u: raise HTTPException(404, "Not found")
    history = await db.login_history.find({"$or": [{"user_id": emp_id}, {"email": u["email"]}]}).sort("at", -1).limit(50).to_list(50)
    return [serialize(h) for h in history]

# ---------- Leads ----------
def lead_health_score(lead: dict) -> int:
    score = 10
    if lead.get("contact_made"): score += 10
    sv = lead.get("site_visit") or {}
    if sv.get("gps_lat"): score += 30
    if sv.get("photos"): score += 20
    if sv.get("requirement_notes"): score += 20
    if lead.get("quotation_count", 0) > 0: score += 10
    return min(score, 100)

@api.get("/leads")
async def list_leads(user: dict = Depends(get_current_user), status: Optional[str] = None,
                     lead_type: Optional[str] = None, assigned: Optional[str] = None):
    q = {}
    if user["role"] in ("field_sales", "store_sales"):
        q["assigned_to"] = user["id"]
    if status: q["status"] = status
    if lead_type: q["lead_type"] = lead_type
    if assigned: q["assigned_to"] = assigned
    leads = await db.leads.find(q).sort("created_at", -1).limit(1000).to_list(1000)
    if not leads:
        return []
    # Batch quotation counts in a single aggregation
    lead_ids = [str(l["_id"]) for l in leads]
    pipeline = [
        {"$match": {"lead_id": {"$in": lead_ids}}},
        {"$group": {"_id": "$lead_id", "count": {"$sum": 1}}},
    ]
    counts_doc = await db.quotations.aggregate(pipeline).to_list(2000)
    counts = {c["_id"]: c["count"] for c in counts_doc}
    result = []
    for l in leads:
        l = serialize(l)
        l["quotation_count"] = counts.get(l["id"], 0)
        l["health_score"] = lead_health_score(l)
        result.append(l)
    return result

@api.post("/leads")
async def create_lead(body: LeadIn, user: dict = Depends(get_current_user)):
    doc = body.model_dump()
    doc["created_by"] = user["id"]
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    doc["updated_at"] = doc["created_at"]
    if not doc.get("assigned_to"):
        doc["assigned_to"] = user["id"]
    res = await db.leads.insert_one(doc)
    lead_id = str(res.inserted_id)
    await db.activities.insert_one({
        "lead_id": lead_id, "activity_type": "lead_created",
        "description": f"Lead created by {user['name']}",
        "user_id": user["id"], "user_name": user["name"],
        "created_at": datetime.now(timezone.utc).isoformat(),
    })
    doc["id"] = lead_id
    return serialize(doc)

@api.get("/leads/{lead_id}")
async def get_lead(lead_id: str, user: dict = Depends(get_current_user)):
    l = await db.leads.find_one({"_id": ObjectId(lead_id)})
    if not l:
        raise HTTPException(404, "Not found")
    l = serialize(l)
    l["quotation_count"] = await db.quotations.count_documents({"lead_id": lead_id})
    l["health_score"] = lead_health_score(l)
    return l

@api.put("/leads/{lead_id}")
async def update_lead(lead_id: str, body: dict, user: dict = Depends(get_current_user)):
    body.pop("_id", None); body.pop("id", None)
    # Closure type required before Won
    if body.get("status") == "Won":
        existing = await db.leads.find_one({"_id": ObjectId(lead_id)})
        closure_type = body.get("closure_type") or (existing or {}).get("closure_type")
        if not closure_type:
            raise HTTPException(400, "closure_type is required before marking a lead as Won")
    body["updated_at"] = datetime.now(timezone.utc).isoformat()
    await db.leads.update_one({"_id": ObjectId(lead_id)}, {"$set": body})
    if "status" in body:
        await db.activities.insert_one({
            "lead_id": lead_id, "activity_type": "status_change",
            "description": f"Status changed to {body['status']}",
            "user_id": user["id"], "user_name": user["name"],
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
    l = await db.leads.find_one({"_id": ObjectId(lead_id)})
    return serialize(l)

@api.delete("/leads/{lead_id}")
async def delete_lead(lead_id: str, user: dict = Depends(require_roles("super_admin", "admin", "sales_manager"))):
    await db.leads.delete_one({"_id": ObjectId(lead_id)})
    return {"ok": True}

@api.post("/leads/{lead_id}/site-visit")
async def add_site_visit(lead_id: str, body: SiteVisitIn, user: dict = Depends(get_current_user)):
    site_data = body.model_dump()
    site_data["executive_id"] = user["id"]
    site_data["executive_name"] = user["name"]
    site_data["timestamp"] = datetime.now(timezone.utc).isoformat()
    await db.leads.update_one(
        {"_id": ObjectId(lead_id)},
        {"$set": {"site_visit": site_data, "updated_at": site_data["timestamp"]}},
    )
    await db.activities.insert_one({
        "lead_id": lead_id, "activity_type": "gps_visit",
        "description": f"GPS-verified site visit by {user['name']}",
        "user_id": user["id"], "user_name": user["name"],
        "created_at": site_data["timestamp"],
    })
    return {"ok": True}

# ---------- Activities ----------
@api.post("/activities")
async def add_activity(body: ActivityIn, user: dict = Depends(get_current_user)):
    doc = body.model_dump()
    doc["user_id"] = user["id"]
    doc["user_name"] = user["name"]
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    res = await db.activities.insert_one(doc)
    if body.activity_type in ("call", "whatsapp", "meeting"):
        await db.leads.update_one({"_id": ObjectId(body.lead_id)}, {"$set": {"contact_made": True}})
    doc["id"] = str(res.inserted_id)
    return serialize(doc)

@api.get("/leads/{lead_id}/activities")
async def lead_activities(lead_id: str, user: dict = Depends(get_current_user)):
    acts = await db.activities.find({"lead_id": lead_id}).sort("created_at", -1).to_list(500)
    return [serialize(a) for a in acts]

# ---------- Tasks ----------
@api.get("/tasks")
async def list_tasks(user: dict = Depends(get_current_user), lead_id: Optional[str] = None):
    q = {}
    if lead_id: q["lead_id"] = lead_id
    if user["role"] in ("field_sales", "store_sales") and not lead_id:
        q["assigned_to"] = user["id"]
    tasks = await db.tasks.find(q).sort("created_at", -1).to_list(500)
    return [serialize(t) for t in tasks]

@api.post("/tasks")
async def create_task(body: TaskIn, user: dict = Depends(get_current_user)):
    doc = body.model_dump()
    doc["created_by"] = user["id"]
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    if not doc.get("assigned_to"):
        doc["assigned_to"] = user["id"]
    res = await db.tasks.insert_one(doc)
    doc["id"] = str(res.inserted_id)
    return serialize(doc)

@api.put("/tasks/{task_id}")
async def update_task(task_id: str, body: dict, user: dict = Depends(get_current_user)):
    body.pop("_id", None); body.pop("id", None)
    await db.tasks.update_one({"_id": ObjectId(task_id)}, {"$set": body})
    t = await db.tasks.find_one({"_id": ObjectId(task_id)})
    return serialize(t)

# ---------- Quotations ----------
def calc_quotation(items: List[dict], extras: dict = None) -> dict:
    extras = extras or {}
    sub = 0; disc = 0; tax = 0
    for it in items:
        line = it["quantity"] * it["unit_price"]
        d = line * it.get("discount_pct", 0) / 100
        taxable = line - d
        t = taxable * it.get("gst_rate", 18) / 100
        sub += line; disc += d; tax += t
    extras_total = sum([extras.get("installation_charge", 0), extras.get("freight_charge", 0),
                        extras.get("amc_charge", 0), extras.get("misc_charge", 0)])
    extras_tax = extras_total * 0.18  # default 18% GST on services
    grand = sub - disc + tax + extras_total + extras_tax
    return {"subtotal": round(sub, 2), "discount": round(disc, 2),
            "tax": round(tax + extras_tax, 2), "extras": round(extras_total, 2),
            "total": round(grand, 2)}

@api.get("/quotations")
async def list_quotations(user: dict = Depends(get_current_user), lead_id: Optional[str] = None):
    q = {}
    if lead_id: q["lead_id"] = lead_id
    quots = await db.quotations.find(q).sort("created_at", -1).to_list(500)
    return [serialize(qq) for qq in quots]

@api.post("/quotations")
async def create_quotation(body: QuotationIn, user: dict = Depends(get_current_user)):
    items = [i.model_dump() for i in body.items]
    extras = {"installation_charge": body.installation_charge, "freight_charge": body.freight_charge,
              "amc_charge": body.amc_charge, "misc_charge": body.misc_charge}
    totals = calc_quotation(items, extras)
    count = await db.quotations.count_documents({})
    qnum = body.quotation_number or f"NMP-Q-{datetime.now().year}-{count+1:05d}"
    doc = {
        "lead_id": body.lead_id,
        "quotation_number": qnum,
        "items": items, "terms": body.terms, "notes": body.notes,
        "totals": totals, "version": 1, "status": body.status or "Draft",
        "installation_charge": body.installation_charge, "freight_charge": body.freight_charge,
        "amc_charge": body.amc_charge, "misc_charge": body.misc_charge,
        "created_by": user["id"], "created_by_name": user["name"],
        "created_at": datetime.now(timezone.utc).isoformat(),
    }
    res = await db.quotations.insert_one(doc)
    await db.activities.insert_one({
        "lead_id": body.lead_id, "activity_type": "quotation_created",
        "description": f"Quotation {qnum} created (Total INR {totals['total']})",
        "user_id": user["id"], "user_name": user["name"],
        "created_at": doc["created_at"],
    })
    doc["id"] = str(res.inserted_id)
    return serialize(doc)

@api.get("/quotations/{qid}")
async def get_quotation(qid: str, user: dict = Depends(get_current_user)):
    q = await db.quotations.find_one({"_id": ObjectId(qid)})
    if not q:
        raise HTTPException(404, "Not found")
    return serialize(q)

@api.put("/quotations/{qid}/status")
async def update_quotation_status(qid: str, body: dict, user: dict = Depends(get_current_user)):
    status = body.get("status")
    if status not in ["Draft", "Sent", "Viewed", "Negotiation", "Approved", "Rejected"]:
        raise HTTPException(400, "Invalid status")
    await db.quotations.update_one({"_id": ObjectId(qid)}, {"$set": {"status": status}})
    q = await db.quotations.find_one({"_id": ObjectId(qid)})
    if q:
        await db.activities.insert_one({
            "lead_id": q["lead_id"], "activity_type": "quotation_status",
            "description": f"Quotation {q['quotation_number']} marked {status}",
            "user_id": user["id"], "user_name": user["name"],
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
    return serialize(q)

# ---------- Follow-ups & Pipeline ----------
@api.get("/follow-ups")
async def follow_ups(user: dict = Depends(get_current_user)):
    q = {"next_follow_up": {"$exists": True, "$ne": ""}}
    if user["role"] in ("field_sales", "store_sales"):
        q["assigned_to"] = user["id"]
    leads = await db.leads.find(q).sort("next_follow_up", 1).to_list(500)
    today = datetime.now(timezone.utc).date().isoformat()
    result = {"today": [], "overdue": [], "upcoming": []}
    for l in leads:
        l = serialize(l)
        d = (l.get("next_follow_up") or "")[:10]
        if not d: continue
        if d < today: result["overdue"].append(l)
        elif d == today: result["today"].append(l)
        else: result["upcoming"].append(l)
    return result

@api.get("/pipeline")
async def pipeline(user: dict = Depends(get_current_user)):
    q = {}
    if user["role"] in ("field_sales", "store_sales"):
        q["assigned_to"] = user["id"]
    stages = ["New", "Contacted", "Site Visit", "Quotation", "Negotiation", "Won", "Lost"]
    out = {}
    for s in stages:
        leads = await db.leads.find({**q, "status": s}).sort("created_at", -1).to_list(200)
        items = [serialize(l) for l in leads]
        out[s] = {
            "count": len(items),
            "value": sum(l.get("expected_deal_value", 0) for l in items),
            "leads": items,
        }
    return out

# ---------- Inventory ----------
@api.get("/inventory")
async def list_inventory(user: dict = Depends(get_current_user), search: Optional[str] = None):
    q = {}
    if search:
        q = {"$or": [{"product_name": {"$regex": search, "$options": "i"}},
                     {"sku": {"$regex": search, "$options": "i"}},
                     {"brand": {"$regex": search, "$options": "i"}}]}
    items = await db.inventory.find(q).limit(500).to_list(500)
    return [serialize(i) for i in items]

@api.post("/inventory")
async def create_inventory(body: InventoryItem, user: dict = Depends(require_roles("super_admin", "admin"))):
    doc = body.model_dump()
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    res = await db.inventory.insert_one(doc)
    doc["id"] = str(res.inserted_id)
    return serialize(doc)

@api.post("/inventory/bulk-import")
async def bulk_import(items: List[InventoryItem], user: dict = Depends(require_roles("super_admin", "admin"))):
    docs = []
    for it in items:
        d = it.model_dump()
        d["created_at"] = datetime.now(timezone.utc).isoformat()
        docs.append(d)
    if docs:
        await db.inventory.insert_many(docs)
    return {"imported": len(docs)}

@api.put("/inventory/{item_id}")
async def update_inventory(item_id: str, body: dict, user: dict = Depends(require_roles("super_admin", "admin"))):
    body.pop("_id", None); body.pop("id", None)
    await db.inventory.update_one({"_id": ObjectId(item_id)}, {"$set": body})
    item = await db.inventory.find_one({"_id": ObjectId(item_id)})
    return serialize(item)

@api.delete("/inventory/{item_id}")
async def delete_inventory(item_id: str, user: dict = Depends(require_roles("super_admin", "admin"))):
    await db.inventory.delete_one({"_id": ObjectId(item_id)})
    return {"ok": True}

# ---------- Password Change ----------
class ChangePasswordIn(BaseModel):
    old_password: str
    new_password: str

@api.post("/auth/change-password")
async def change_password(body: ChangePasswordIn, user: dict = Depends(get_current_user)):
    u = await db.users.find_one({"_id": ObjectId(user["id"])})
    if not verify_password(body.old_password, u["password_hash"]):
        raise HTTPException(400, "Current password is incorrect")
    if len(body.new_password) < 6:
        raise HTTPException(400, "New password must be at least 6 characters")
    await db.users.update_one(
        {"_id": ObjectId(user["id"])},
        {"$set": {"password_hash": hash_password(body.new_password)}, "$inc": {"token_version": 1}},
    )
    return {"ok": True}

# ---------- GPS Check-in / Check-out ----------
class CheckInIn(BaseModel):
    lead_id: Optional[str] = ""
    gps_lat: float
    gps_lng: float
    gps_accuracy: Optional[float] = 0
    notes: Optional[str] = ""

class CheckOutIn(BaseModel):
    check_in_id: str
    gps_lat: float
    gps_lng: float
    distance_km: Optional[float] = 0
    duration_minutes: Optional[float] = 0
    signature_data: Optional[str] = ""  # base64 image
    customer_name: Optional[str] = ""
    summary: Optional[str] = ""

@api.post("/gps/check-in")
async def gps_check_in(body: CheckInIn, user: dict = Depends(get_current_user)):
    doc = body.model_dump()
    doc["user_id"] = user["id"]
    doc["user_name"] = user["name"]
    doc["check_in_at"] = datetime.now(timezone.utc).isoformat()
    doc["status"] = "Open"
    res = await db.gps_visits.insert_one(doc)
    if body.lead_id:
        await db.activities.insert_one({
            "lead_id": body.lead_id, "activity_type": "gps_check_in",
            "description": f"GPS Check-in by {user['name']} ({body.gps_lat:.5f}, {body.gps_lng:.5f})",
            "user_id": user["id"], "user_name": user["name"],
            "created_at": doc["check_in_at"],
        })
    doc["id"] = str(res.inserted_id)
    return serialize(doc)

@api.post("/gps/check-out")
async def gps_check_out(body: CheckOutIn, user: dict = Depends(get_current_user)):
    visit = await db.gps_visits.find_one({"_id": ObjectId(body.check_in_id), "user_id": user["id"]})
    if not visit:
        raise HTTPException(404, "Check-in not found or not yours")
    update = body.model_dump()
    update["check_out_at"] = datetime.now(timezone.utc).isoformat()
    update["status"] = "Closed"
    await db.gps_visits.update_one({"_id": ObjectId(body.check_in_id)}, {"$set": update})
    if visit.get("lead_id"):
        await db.activities.insert_one({
            "lead_id": visit["lead_id"], "activity_type": "gps_check_out",
            "description": f"Visit closed by {user['name']} (distance {body.distance_km}km, signed by {body.customer_name or 'customer'})",
            "user_id": user["id"], "user_name": user["name"],
            "created_at": update["check_out_at"],
        })
    visit.update(update)
    return serialize(visit)

@api.get("/gps/visits")
async def list_visits(user: dict = Depends(get_current_user), lead_id: Optional[str] = None):
    q = {}
    if lead_id: q["lead_id"] = lead_id
    if user["role"] in ("field_sales", "store_sales"):
        q["user_id"] = user["id"]
    visits = await db.gps_visits.find(q).sort("check_in_at", -1).to_list(200)
    return [serialize(v) for v in visits]

# ---------- Export (Excel) ----------
@api.get("/export/excel")
async def export_excel(user: dict = Depends(require_roles("super_admin"))):
    from openpyxl import Workbook
    from io import BytesIO
    wb = Workbook()
    # Leads
    ws = wb.active
    ws.title = "Leads"
    headers = ["lead_name","company_name","contact_person","mobile","email","city","state","lead_source","lead_type","status","priority","expected_deal_value","probability","next_follow_up","competitor","created_at"]
    ws.append([h.replace("_"," ").title() for h in headers])
    async for l in db.leads.find({}, projection={h: 1 for h in headers}).limit(10000):
        ws.append([str(l.get(h, "") or "") for h in headers])
    # Quotations
    ws2 = wb.create_sheet("Quotations")
    q_headers = ["quotation_number","lead_id","status","subtotal","discount","tax","total","items_count","created_by_name","created_at"]
    ws2.append([h.replace("_"," ").title() for h in q_headers])
    q_proj = {"quotation_number": 1, "lead_id": 1, "status": 1, "totals": 1, "items": 1, "created_by_name": 1, "created_at": 1}
    async for q in db.quotations.find({}, projection=q_proj).limit(10000):
        totals = q.get("totals", {})
        row = [q.get("quotation_number",""), q.get("lead_id",""), q.get("status",""),
               totals.get("subtotal",0), totals.get("discount",0), totals.get("tax",0), totals.get("total",0),
               len(q.get("items", [])), q.get("created_by_name",""), str(q.get("created_at",""))]
        ws2.append(row)
    # Inventory
    ws3 = wb.create_sheet("Inventory")
    i_headers = ["sku","product_name","brand","category","cost_price","selling_price","mrp","gst_rate","stock","warehouse"]
    ws3.append([h.replace("_"," ").title() for h in i_headers])
    async for it in db.inventory.find({}, projection={h: 1 for h in i_headers}).limit(10000):
        ws3.append([str(it.get(h, "") or "") for h in i_headers])
    # Employees
    ws4 = wb.create_sheet("Employees")
    e_headers = ["name","email","mobile","role","department","designation","territory","status","created_at"]
    ws4.append([h.replace("_"," ").title() for h in e_headers])
    async for e in db.users.find({}, projection={h: 1 for h in e_headers}).limit(10000):
        ws4.append([str(e.get(h, "") or "") for h in e_headers])
    # Tasks
    ws5 = wb.create_sheet("Tasks")
    t_headers = ["title","task_type","status","lead_id","assigned_to","due_date","created_at"]
    ws5.append([h.replace("_"," ").title() for h in t_headers])
    async for t in db.tasks.find({}, projection={h: 1 for h in t_headers}).limit(10000):
        ws5.append([str(t.get(h, "") or "") for h in t_headers])
    # GPS Visits
    ws6 = wb.create_sheet("GPS Visits")
    v_headers = ["user_name","lead_id","gps_lat","gps_lng","check_in_at","check_out_at","distance_km","customer_name","status"]
    ws6.append([h.replace("_"," ").title() for h in v_headers])
    async for v in db.gps_visits.find({}, projection={h: 1 for h in v_headers}).limit(10000):
        ws6.append([str(v.get(h, "") or "") for h in v_headers])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    filename = f"NMP_SalesOS_Export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{filename}"'})

# ---------- Dashboard ----------
@api.get("/dashboard/kpis")
async def dashboard_kpis(user: dict = Depends(get_current_user)):
    q = {}
    if user["role"] in ("field_sales", "store_sales"):
        q["assigned_to"] = user["id"]
    total = await db.leads.count_documents(q)
    open_q = {**q, "status": {"$nin": ["Won", "Lost"]}}
    open_leads = await db.leads.count_documents(open_q)
    won = await db.leads.count_documents({**q, "status": "Won"})
    lost = await db.leads.count_documents({**q, "status": "Lost"})
    gps_verified = await db.leads.count_documents({**q, "site_visit.gps_lat": {"$exists": True}})
    visit_pending = await db.leads.count_documents({**q, "visit_requirement": "Mandatory",
                                                    "site_visit.gps_lat": {"$exists": False}})
    quots_q = {}
    if user["role"] in ("field_sales", "store_sales"):
        quots_q["created_by"] = user["id"]
    quotations_sent = await db.quotations.count_documents(quots_q)
    won_leads = await db.leads.find({**q, "status": "Won"}).to_list(1000)
    revenue = sum(l.get("expected_deal_value", 0) for l in won_leads)
    conv_rate = round((won / total * 100), 1) if total else 0
    pipeline = await db.leads.aggregate([
        {"$match": q}, {"$group": {"_id": "$status", "count": {"$sum": 1}}},
    ]).to_list(50)
    status_breakdown = {p["_id"]: p["count"] for p in pipeline}
    src_pipe = await db.leads.aggregate([
        {"$match": q}, {"$group": {"_id": "$lead_source", "count": {"$sum": 1}}},
    ]).to_list(50)
    source_breakdown = [{"source": s["_id"], "count": s["count"]} for s in src_pipe]
    # Pipeline value (open leads expected value)
    open_leads_docs = await db.leads.find(open_q).to_list(2000)
    pipeline_value = sum(l.get("expected_deal_value", 0) for l in open_leads_docs)
    # Follow ups
    today_iso = datetime.now(timezone.utc).date().isoformat()
    fu_q = {"next_follow_up": {"$exists": True, "$ne": ""}}
    if user["role"] in ("field_sales", "store_sales"):
        fu_q["assigned_to"] = user["id"]
    fu_leads = await db.leads.find(fu_q).to_list(2000)
    today_fu = sum(1 for l in fu_leads if (l.get("next_follow_up") or "")[:10] == today_iso)
    overdue_fu = sum(1 for l in fu_leads if (l.get("next_follow_up") or "")[:10] and (l.get("next_follow_up") or "")[:10] < today_iso)
    return {"total_leads": total, "open_leads": open_leads, "deals_won": won, "deals_lost": lost,
            "gps_verified": gps_verified, "visit_pending": visit_pending,
            "quotations_sent": quotations_sent, "revenue": revenue,
            "conversion_rate": conv_rate, "status_breakdown": status_breakdown,
            "source_breakdown": source_breakdown,
            "pipeline_value": pipeline_value,
            "today_follow_ups": today_fu, "overdue_follow_ups": overdue_fu,
            "quotation_approval_rate": await quotation_approval_rate(quots_q)}

async def quotation_approval_rate(quots_q: dict) -> float:
    sent_or_more = await db.quotations.count_documents({**quots_q, "status": {"$in": ["Sent", "Viewed", "Negotiation", "Approved", "Rejected"]}})
    if not sent_or_more:
        return 0.0
    approved = await db.quotations.count_documents({**quots_q, "status": "Approved"})
    return round(approved / sent_or_more * 100, 1)

# ---------- File Upload ----------
@api.post("/files/upload")
async def upload_file(file: UploadFile = File(...), user: dict = Depends(get_current_user)):
    content = await file.read()
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(400, "File too large (max 5MB)")
    encoded = base64.b64encode(content).decode("utf-8")
    doc = {"filename": file.filename, "content_type": file.content_type,
           "data": encoded, "size": len(content),
           "uploaded_by": user["id"],
           "uploaded_at": datetime.now(timezone.utc).isoformat()}
    res = await db.files.insert_one(doc)
    fid = str(res.inserted_id)
    return {"id": fid, "url": f"/api/files/{fid}", "filename": file.filename,
            "content_type": file.content_type}

@api.get("/files/{file_id}")
async def get_file(file_id: str):
    doc = await db.files.find_one({"_id": ObjectId(file_id)})
    if not doc:
        raise HTTPException(404, "Not found")
    content = base64.b64decode(doc["data"])
    return Response(content=content, media_type=doc.get("content_type", "application/octet-stream"))

# ---------- Knowledge / Playbooks / Training (DB-backed with seed defaults) ----------
DEFAULT_PLAYBOOKS = [
    {"id": "church", "category": "Church Sales",
     "discovery_questions": ["Seating capacity?", "Reverb/echo issues?", "Existing PA system?", "Worship style?", "Live streaming required?"],
     "objections": [{"o": "Too expensive", "r": "Show 10-year TCO vs cheap systems."}, {"o": "We will think about it", "r": "Offer free acoustic survey."}],
     "products": ["Line array speakers", "Wireless mics (Shure/Sennheiser)", "Digital mixer", "In-ear monitoring", "Acoustic treatment"],
     "checklist": ["Site survey done", "GPS verified", "Acoustic photos", "Stakeholder list", "Quotation v1 sent"]},
    {"id": "school", "category": "School Sales",
     "discovery_questions": ["Auditorium size?", "Assembly type?", "Music classroom needs?", "Existing PA?", "Budget cycle?"],
     "objections": [{"o": "Government tender required", "r": "We provide GeM-compliant quotes."}],
     "products": ["Column speakers", "Lapel mic systems", "Portable PA", "Conference mics"],
     "checklist": ["Principal meeting", "Site visit + photos", "Tender format", "Proposal submitted"]},
    {"id": "studio", "category": "Studio Sales",
     "discovery_questions": ["Genre focus?", "Tracks recorded?", "DAW preference?", "Monitoring?", "Treatment?"],
     "objections": [{"o": "I can buy online", "r": "We offer installation + calibration + warranty."}],
     "products": ["Audio interfaces", "Studio monitors", "Acoustic treatment", "Condenser mics", "Headphones"],
     "checklist": ["Studio walkthrough", "Acoustic photos", "Signal chain", "Demo session"]},
    {"id": "corporate_av", "category": "Corporate AV",
     "discovery_questions": ["Meeting rooms?", "VC platform?", "BYOD or fixed?", "Control system?"],
     "objections": [{"o": "IT will handle it", "r": "Offer joint demo with IT head."}],
     "products": ["Ceiling mics", "Video bars", "Display screens", "Control automation"],
     "checklist": ["Floor plan", "IT alignment", "PoC done", "Maintenance AMC quoted"]},
    {"id": "event", "category": "Event Companies",
     "discovery_questions": ["Avg event size?", "Rental fleet gaps?", "Tour requirements?"],
     "objections": [{"o": "Cash flow tight", "r": "Easy EMI / lease finance."}],
     "products": ["Line arrays", "Subwoofers", "Wireless systems", "Lighting controllers"],
     "checklist": ["Warehouse demo", "Trial event", "Service contract", "Spare kit"]},
    {"id": "government", "category": "Government Projects",
     "discovery_questions": ["Tender ID?", "Compliance docs?", "Site location?", "Timeline?"],
     "objections": [{"o": "L1 only", "r": "Highlight 5-year service inclusion."}],
     "products": ["IP-rated speakers", "PA systems", "Conference systems", "Translation"],
     "checklist": ["Tender review", "Site visit", "Bid submitted", "Tech approval"]},
]

DEFAULT_KNOWLEDGE = [
    {"category": "Product Guide", "title": "Choosing the Right Line Array",
     "body": "Line arrays project sound over long throws with even coverage. Use for auditoriums over 300 seats, churches with high ceilings, or outdoor events.", "published": True},
    {"category": "Brand Guide", "title": "Shure SM58 vs Sennheiser e835",
     "body": "SM58: industry-standard, very durable. e835: brighter top end, lower handling noise. Recommend SM58 for live worship; e835 for spoken-word.", "published": True},
    {"category": "FAQ", "title": "Treatment vs Soundproofing?",
     "body": "Treatment improves sound INSIDE a room. Soundproofing stops sound from passing THROUGH walls. Churches need treatment; studios often need both.", "published": True},
    {"category": "Troubleshooting", "title": "Feedback during worship",
     "body": "Check mic placement vs monitors. Engage HPF below 80Hz. Use EQ to notch ringing freq. Consider in-ear monitors.", "published": True},
    {"category": "Competitor", "title": "Why NMP vs Online Stores",
     "body": "Online stores cannot do site surveys, acoustic measurement, installation, training, or AMC. NMP provides full lifecycle support.", "published": True},
    {"category": "Sales Talk", "title": "Opening a Church Conversation",
     "body": "Ask: How did Sunday service feel? Could everyone hear clearly? When was the last system upgrade?", "published": True},
    {"category": "Installation", "title": "Auditorium Install Checklist",
     "body": "Cable runs, conduit, rigging certification, power isolation, equipment rack, signal flow diagram, commissioning report, customer training.", "published": True},
]

DEFAULT_TRAINING = [
    {"title": "NMP Sales Process 101", "category": "Sales", "duration": "20 min",
     "description": "Lead-centric selling. Site visit discipline, GPS, photos, follow-up cadence.",
     "video_url": "https://www.youtube.com/embed/dQw4w9WgXcQ", "pdf_url": "",
     "video_type": "Training Video",
     "quiz": [
         {"q": "What is the single source of truth in NMP?", "options": ["Quotation", "Lead", "Task", "Photo"], "answer": 1},
         {"q": "When is GPS verification mandatory?", "options": ["Never", "For Product Sale only", "When Visit Requirement is Mandatory", "Optional always"], "answer": 2},
         {"q": "Maximum lead health score?", "options": ["50", "75", "100", "120"], "answer": 2},
     ]},
    {"title": "Shure SM58 Product Demo", "category": "Product", "duration": "8 min",
     "description": "Hands-on demo of the industry-standard dynamic vocal mic — features, benefits, sales talking points.",
     "video_url": "https://www.youtube.com/embed/dQw4w9WgXcQ", "pdf_url": "",
     "video_type": "Product Demo",
     "quiz": [
         {"q": "SM58 pickup pattern?", "options": ["Omni", "Cardioid", "Figure-8", "Shotgun"], "answer": 1},
         {"q": "Best application for SM58?", "options": ["Studio recording", "Live vocals", "Boundary mic", "Lavalier"], "answer": 1},
     ]},
    {"title": "Yamaha DXR Series Demo", "category": "Product", "duration": "12 min",
     "description": "Yamaha DXR active speaker series — coverage, SPL, daisy-chain configurations.",
     "video_url": "https://www.youtube.com/embed/dQw4w9WgXcQ", "pdf_url": "",
     "video_type": "Product Demo",
     "quiz": [
         {"q": "DXR10 wattage class?", "options": ["350W", "700W", "1100W", "1300W"], "answer": 2},
         {"q": "DXR is active or passive?", "options": ["Active (powered)", "Passive (unpowered)"], "answer": 0},
     ]},
    {"title": "Auditorium Install SOP", "category": "Technical", "duration": "45 min",
     "description": "Step-by-step installation SOP for auditoriums — cable runs, rigging, commissioning, customer handover.",
     "video_url": "https://www.youtube.com/embed/dQw4w9WgXcQ", "pdf_url": "",
     "video_type": "SOP Video",
     "quiz": [
         {"q": "Rigging certification is:", "options": ["Optional", "Mandatory for flown speakers", "For mics only", "Customer choice"], "answer": 1},
         {"q": "Final commissioning includes:", "options": ["EQ calibration", "Signal flow doc", "Customer training", "All of the above"], "answer": 3},
     ]},
    {"title": "Site Survey SOP", "category": "Technical", "duration": "25 min",
     "description": "Field SOP for conducting a complete site survey — photos, measurements, requirement capture, GPS verification.",
     "video_url": "https://www.youtube.com/embed/dQw4w9WgXcQ", "pdf_url": "",
     "video_type": "SOP Video",
     "quiz": [
         {"q": "Before quotation, what must be GPS verified?", "options": ["Office address", "Site location", "Vendor", "Bank"], "answer": 1},
         {"q": "Minimum site photos required?", "options": ["0", "1", "3", "5+"], "answer": 3},
     ]},
    {"title": "Quotation Writing Masterclass", "category": "Sales", "duration": "15 min",
     "description": "Structuring quotes, discount discipline, GST, payment terms, closing techniques.",
     "video_url": "https://www.youtube.com/embed/dQw4w9WgXcQ", "pdf_url": "",
     "video_type": "Training Video",
     "quiz": [
         {"q": "Standard GST on pro-audio products in India?", "options": ["5%", "12%", "18%", "28%"], "answer": 2},
         {"q": "NMP default advance is:", "options": ["10%", "25%", "50%", "100%"], "answer": 2},
     ]},
]

# Playbooks CRUD
class PlaybookIn(BaseModel):
    category: str
    discovery_questions: List[str] = []
    objections: List[dict] = []
    products: List[str] = []
    checklist: List[str] = []
    archived: bool = False

@api.get("/playbooks")
async def get_playbooks(user: dict = Depends(get_current_user)):
    pbs = await db.playbooks.find({"archived": {"$ne": True}}).to_list(200)
    return [serialize(p) for p in pbs]

@api.post("/playbooks")
async def create_playbook(body: PlaybookIn, user: dict = Depends(require_roles("super_admin", "admin"))):
    doc = body.model_dump()
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    res = await db.playbooks.insert_one(doc)
    doc["id"] = str(res.inserted_id)
    return serialize(doc)

@api.put("/playbooks/{pb_id}")
async def update_playbook(pb_id: str, body: dict, user: dict = Depends(require_roles("super_admin", "admin"))):
    body.pop("_id", None); body.pop("id", None)
    await db.playbooks.update_one({"_id": ObjectId(pb_id)}, {"$set": body})
    p = await db.playbooks.find_one({"_id": ObjectId(pb_id)})
    return serialize(p)

@api.delete("/playbooks/{pb_id}")
async def delete_playbook(pb_id: str, user: dict = Depends(require_roles("super_admin", "admin"))):
    await db.playbooks.delete_one({"_id": ObjectId(pb_id)})
    return {"ok": True}

# Knowledge CRUD
class KnowledgeIn(BaseModel):
    category: str
    title: str
    body: str
    published: bool = True

@api.get("/knowledge")
async def get_knowledge(user: dict = Depends(get_current_user)):
    items = await db.knowledge.find({"published": True}).sort("created_at", -1).to_list(500)
    return [serialize(i) for i in items]

@api.post("/knowledge")
async def create_knowledge(body: KnowledgeIn, user: dict = Depends(require_roles("super_admin", "admin"))):
    doc = body.model_dump()
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    res = await db.knowledge.insert_one(doc)
    doc["id"] = str(res.inserted_id)
    return serialize(doc)

@api.put("/knowledge/{k_id}")
async def update_knowledge(k_id: str, body: dict, user: dict = Depends(require_roles("super_admin", "admin"))):
    body.pop("_id", None); body.pop("id", None)
    await db.knowledge.update_one({"_id": ObjectId(k_id)}, {"$set": body})
    k = await db.knowledge.find_one({"_id": ObjectId(k_id)})
    return serialize(k)

@api.delete("/knowledge/{k_id}")
async def delete_knowledge(k_id: str, user: dict = Depends(require_roles("super_admin", "admin"))):
    await db.knowledge.delete_one({"_id": ObjectId(k_id)})
    return {"ok": True}

# Training (LMS)
class TrainingIn(BaseModel):
    title: str
    category: str
    duration: str = ""
    description: str = ""
    video_url: str = ""
    pdf_url: str = ""
    video_type: str = "Training Video"  # Training Video / Product Demo / SOP Video
    quiz: List[dict] = []

@api.get("/training")
async def get_training(user: dict = Depends(get_current_user)):
    items = await db.training.find().sort("created_at", -1).to_list(200)
    out = []
    for t in items:
        t = serialize(t)
        p = await db.training_progress.find_one({"training_id": t["id"], "user_id": user["id"]})
        t["progress"] = serialize(p) if p else {"status": "Not Started", "score": 0, "video_watched": False}
        out.append(t)
    return out

@api.post("/training")
async def create_training(body: TrainingIn, user: dict = Depends(require_roles("super_admin", "admin"))):
    doc = body.model_dump()
    doc["created_at"] = datetime.now(timezone.utc).isoformat()
    res = await db.training.insert_one(doc)
    doc["id"] = str(res.inserted_id)
    return serialize(doc)

@api.put("/training/{t_id}")
async def update_training(t_id: str, body: dict, user: dict = Depends(require_roles("super_admin", "admin"))):
    body.pop("_id", None); body.pop("id", None)
    await db.training.update_one({"_id": ObjectId(t_id)}, {"$set": body})
    t = await db.training.find_one({"_id": ObjectId(t_id)})
    return serialize(t)

@api.delete("/training/{t_id}")
async def delete_training(t_id: str, user: dict = Depends(require_roles("super_admin", "admin"))):
    await db.training.delete_one({"_id": ObjectId(t_id)})
    await db.training_progress.delete_many({"training_id": t_id})
    return {"ok": True}

@api.post("/training/{t_id}/mark-watched")
async def mark_watched(t_id: str, user: dict = Depends(get_current_user)):
    await db.training_progress.update_one(
        {"training_id": t_id, "user_id": user["id"]},
        {"$set": {"training_id": t_id, "user_id": user["id"], "user_name": user["name"],
                  "video_watched": True, "status": "In Progress",
                  "watched_at": datetime.now(timezone.utc).isoformat()}},
        upsert=True,
    )
    return {"ok": True}

@api.post("/training/{t_id}/submit-quiz")
async def submit_quiz(t_id: str, body: dict, user: dict = Depends(get_current_user)):
    t = await db.training.find_one({"_id": ObjectId(t_id)})
    if not t: raise HTTPException(404, "Not found")
    answers = body.get("answers", [])
    quiz = t.get("quiz", [])
    correct = sum(1 for i, q in enumerate(quiz) if i < len(answers) and answers[i] == q.get("answer"))
    total = len(quiz)
    score = round(correct / total * 100, 1) if total else 0
    passed = score >= 70
    cert_number = ""
    if passed:
        cert_number = f"NMP-CERT-{datetime.now().year}-{secrets.token_hex(3).upper()}"
    progress = {
        "training_id": t_id, "user_id": user["id"], "user_name": user["name"],
        "status": "Certified" if passed else "Failed",
        "score": score, "answers": answers,
        "completed_at": datetime.now(timezone.utc).isoformat(),
        "certificate_number": cert_number,
    }
    await db.training_progress.update_one(
        {"training_id": t_id, "user_id": user["id"]},
        {"$set": progress}, upsert=True
    )
    return {"score": score, "passed": passed, "certificate_number": cert_number,
            "training_title": t.get("title")}

@api.get("/training/{t_id}/certificate")
async def get_certificate(t_id: str, user: dict = Depends(get_current_user)):
    p = await db.training_progress.find_one({"training_id": t_id, "user_id": user["id"]})
    if not p or p.get("status") != "Certified":
        raise HTTPException(404, "No certificate found")
    t = await db.training.find_one({"_id": ObjectId(t_id)})
    return {"certificate_number": p.get("certificate_number"),
            "user_name": user["name"], "user_id": user["id"],
            "training_title": t.get("title"),
            "score": p.get("score"),
            "issue_date": p.get("completed_at", "").split("T")[0]}

# Stale quotation nudge
@api.post("/quotations/run-stale-nudge")
async def run_stale_nudge(user: dict = Depends(require_roles("super_admin", "admin", "sales_manager"))):
    cutoff = (datetime.now(timezone.utc) - timedelta(days=3)).isoformat()
    stale = await db.quotations.find({"status": "Sent", "created_at": {"$lt": cutoff}}).to_list(500)
    created = 0
    for q in stale:
        # Skip if a nudge task already exists
        existing = await db.tasks.find_one({"lead_id": q["lead_id"],
                                            "title": {"$regex": f"^Follow-up: {q['quotation_number']}"}})
        if existing: continue
        await db.tasks.insert_one({
            "lead_id": q["lead_id"],
            "title": f"Follow-up: {q['quotation_number']} pending >3 days",
            "description": f"Quotation {q['quotation_number']} for INR {q.get('totals',{}).get('total',0)} has been Sent but not actioned.",
            "task_type": "Follow-Up",
            "assigned_to": q.get("created_by", ""),
            "status": "Open",
            "created_by": user["id"],
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
        created += 1
    return {"stale_count": len(stale), "tasks_created": created}

# ---------- Wiring ----------
app.include_router(api)

origins = [o.strip() for o in os.environ.get('CORS_ORIGINS', '*').split(',') if o.strip()]
app.add_middleware(
    CORSMiddleware,
    allow_credentials=True,
    allow_origins=origins if origins != ['*'] else ['*'],
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.on_event("startup")
async def startup():
    await db.users.create_index("email", unique=True)
    await db.leads.create_index("assigned_to")
    await db.leads.create_index("status")
    await db.login_history.create_index("user_id")
    if await db.users.count_documents({"role": "super_admin"}) == 0:
        await db.users.insert_one({
            "email": ADMIN_EMAIL.lower(),
            "password_hash": hash_password(ADMIN_PASSWORD),
            "name": "Super Admin",
            "role": "super_admin",
            "status": "Active",
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
        logger.info(f"Seeded super admin: {ADMIN_EMAIL}")
    # Seed playbooks
    if await db.playbooks.count_documents({}) == 0:
        for pb in DEFAULT_PLAYBOOKS:
            d = dict(pb)
            d.pop("id", None)
            d["created_at"] = datetime.now(timezone.utc).isoformat()
            await db.playbooks.insert_one(d)
        logger.info("Seeded default playbooks")
    # Seed knowledge
    if await db.knowledge.count_documents({}) == 0:
        for k in DEFAULT_KNOWLEDGE:
            d = dict(k)
            d["created_at"] = datetime.now(timezone.utc).isoformat()
            await db.knowledge.insert_one(d)
        logger.info("Seeded default knowledge")
    # Seed training
    if await db.training.count_documents({}) == 0:
        for t in DEFAULT_TRAINING:
            d = dict(t)
            d["created_at"] = datetime.now(timezone.utc).isoformat()
            await db.training.insert_one(d)
        logger.info("Seeded default training")

@app.on_event("shutdown")
async def shutdown():
    client.close()
